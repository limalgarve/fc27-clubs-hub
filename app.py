from __future__ import annotations
import io, math, os, secrets, hashlib, hmac
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
import pandas as pd
import psycopg
from psycopg.rows import dict_row
from fastapi import FastAPI, HTTPException, Query, Request, Response
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from fc27_api import FC27API, FC27APIError

BASE=Path(__file__).resolve().parent
DATABASE_URL=os.getenv("DATABASE_URL","").strip()
ADMIN_TOKEN=os.getenv("ADMIN_TOKEN","").strip()
MATCH_TYPES={"leagueMatch":"Liga","friendlyMatch":"Amistoso","playoffMatch":"Playoff"}
api=FC27API(timezone="America/Sao_Paulo",timeout=20)
app=FastAPI(title="FC27 Clubs Hub")
app.mount("/static",StaticFiles(directory=BASE/"static"),name="static")
templates=Jinja2Templates(directory=BASE/"templates")

def clean(v):
    if v is None:return None
    if isinstance(v,float) and math.isnan(v):return None
    if isinstance(v,pd.Timestamp):
        if v.tzinfo is not None:v=v.tz_convert("America/Sao_Paulo").tz_localize(None)
        return v.isoformat(sep=" ",timespec="seconds")
    if hasattr(v,"item"):
        try:return v.item()
        except Exception:pass
    return v

def records(df):
    if df is None or getattr(df,"empty",True):return []
    return [{k:clean(v) for k,v in row.items()} for row in df.to_dict("records")]

@contextmanager
def conn():
    if not DATABASE_URL: raise RuntimeError("DATABASE_URL não configurada.")
    c=psycopg.connect(DATABASE_URL,row_factory=dict_row)
    try:
        yield c;c.commit()
    except Exception:
        c.rollback();raise
    finally:c.close()

def qone(c,sql,args=()):
    with c.cursor() as cur:cur.execute(sql,args);return cur.fetchone()
def qall(c,sql,args=()):
    with c.cursor() as cur:cur.execute(sql,args);return cur.fetchall()
def execq(c,sql,args=()):
    with c.cursor() as cur:cur.execute(sql,args)

def init_db():
    with conn() as c:
        for sql in [
        """CREATE TABLE IF NOT EXISTS users(id BIGSERIAL PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,is_admin BOOLEAN NOT NULL DEFAULT FALSE,active BOOLEAN NOT NULL DEFAULT TRUE,created_at TIMESTAMP NOT NULL DEFAULT NOW())""",
        """CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,expires_at TIMESTAMP NOT NULL DEFAULT (NOW()+INTERVAL '30 days'))""",
        """CREATE TABLE IF NOT EXISTS user_clubs(user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,club_id TEXT NOT NULL,PRIMARY KEY(user_id,club_id))""",
        """CREATE TABLE IF NOT EXISTS clubs(club_id TEXT PRIMARY KEY,name TEXT,stadium TEXT,region_id TEXT,team_id TEXT,games_played INTEGER,wins INTEGER,ties INTEGER,losses INTEGER,goals INTEGER,goals_against INTEGER,skill_rating DOUBLE PRECISION,win_streak INTEGER,unbeaten_streak INTEGER,promotions INTEGER,relegations INTEGER,best_division INTEGER,updated_at TIMESTAMP)""",
        """CREATE TABLE IF NOT EXISTS members(club_id TEXT,name TEXT,player_name TEXT,position TEXT,overall DOUBLE PRECISION,games_played INTEGER,win_rate DOUBLE PRECISION,goals INTEGER,assists INTEGER,average_rating DOUBLE PRECISION,motm INTEGER,shot_success_rate DOUBLE PRECISION,passes_made INTEGER,pass_success_rate DOUBLE PRECISION,tackles_made INTEGER,tackle_success_rate DOUBLE PRECISION,clean_sheets_def INTEGER,clean_sheets_gk INTEGER,red_cards INTEGER,updated_at TIMESTAMP,PRIMARY KEY(club_id,name))""",
        """CREATE TABLE IF NOT EXISTS matches(club_id TEXT,match_id TEXT,match_type TEXT,timestamp TIMESTAMP,club_name TEXT,opponent_id TEXT,opponent_name TEXT,goals INTEGER,goals_against INTEGER,result TEXT,dnf INTEGER,PRIMARY KEY(club_id,match_id))""",
        """CREATE TABLE IF NOT EXISTS player_matches(club_id TEXT,match_id TEXT,match_type TEXT,timestamp TIMESTAMP,player_id TEXT,name TEXT,position TEXT,rating DOUBLE PRECISION,goals INTEGER,assists INTEGER,shots INTEGER,passes_made INTEGER,pass_attempts INTEGER,tackles_made INTEGER,tackle_attempts INTEGER,saves INTEGER,clean_sheet INTEGER,motm INTEGER,red_cards INTEGER,seconds_played INTEGER,PRIMARY KEY(club_id,match_id,player_id,name))"""]:
            execq(c,sql)
        if ADMIN_TOKEN:
            r=qone(c,"SELECT id FROM users WHERE username=%s",("admin",))
            if not r: execq(c,"INSERT INTO users(username,password_hash,is_admin) VALUES(%s,%s,TRUE)",("admin",hash_password(ADMIN_TOKEN)))
            else: execq(c,"UPDATE users SET password_hash=%s,is_admin=TRUE,active=TRUE WHERE username=%s",(hash_password(ADMIN_TOKEN),"admin"))

def hash_password(p):
    salt=secrets.token_bytes(16);dk=hashlib.pbkdf2_hmac("sha256",p.encode(),salt,200000)
    return salt.hex()+"$"+dk.hex()
def verify_password(p,stored):
    try:
        s,h=stored.split("$",1);dk=hashlib.pbkdf2_hmac("sha256",p.encode(),bytes.fromhex(s),200000)
        return hmac.compare_digest(dk.hex(),h)
    except:return False
def current_user(request:Request):
    tok=request.cookies.get("fc27_session")
    if not tok:return None
    with conn() as c:return qone(c,"SELECT u.id,u.username,u.is_admin FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=%s AND s.expires_at>NOW() AND u.active=TRUE",(tok,))
def require_user(request):
    u=current_user(request)
    if not u:raise HTTPException(401,"Faça login para continuar.")
    return u
def require_admin(request):
    u=require_user(request)
    if not u["is_admin"]:raise HTTPException(403,"Acesso somente para administrador.")
    return u
def can_access(c,u,club_id):
    if u["is_admin"]:return True
    return bool(qone(c,"SELECT 1 FROM user_clubs WHERE user_id=%s AND club_id=%s",(u["id"],club_id)))
def assert_club(request,club_id):
    u=require_user(request)
    with conn() as c:
        if not can_access(c,u,club_id):raise HTTPException(403,"Este clube não pertence à sua conta.")
    return u

def safe_matches(club_id,match_type):
    raw=api._get_matches(str(club_id),match_type,10) or [];out=[]
    for m in raw:
        if not isinstance(m,dict):continue
        clubs=m.get("clubs") or {};us=clubs.get(str(club_id)) or {}
        opp_id=None;them={}
        for oid,oval in clubs.items():
            if str(oid)!=str(club_id):opp_id=str(oid);them=oval if isinstance(oval,dict) else {};break
        def i(v):
            try:return int(float(v or 0))
            except:return 0
        gp=i(us.get("goals"));gc=i(us.get("goalsAgainst"))
        result="win" if str(us.get("wins"))=="1" or gp>gc else "loss" if str(us.get("losses"))=="1" or gp<gc else "draw"
        ud=us.get("details") or {};td=them.get("details") or {}
        try:t=pd.to_datetime(m.get("timestamp"),unit="s",utc=True).tz_convert("America/Sao_Paulo").tz_localize(None).to_pydatetime()
        except:
            try:t=pd.to_datetime(m.get("timestamp"),utc=True).tz_convert("America/Sao_Paulo").tz_localize(None).to_pydatetime()
            except:t=datetime.now()
        out.append({"matchId":str(m.get("matchId") or ""),"timestamp":t,"clubName":ud.get("name"),"opponentId":opp_id,"opponentName":td.get("name") or "Adversário","goals":gp,"goalsAgainst":gc,"result":result,"dnf":str(us.get("winnerByDnf"))=="1" or str(them.get("winnerByDnf"))=="1","_raw":m})
    return out

def update_club(club_name,user_id):
    club_id=str(api.find_club_id(club_name));details=records(api.get_club_details(club_id));overall=records(api.get_club_overall_stats(club_id));members=records(api.get_member_stats(club_id))
    d=details[0] if details else {};o=overall[0] if overall else {};official=d.get("clubName") or club_name;now=datetime.now()
    with conn() as c:
        execq(c,"""INSERT INTO clubs VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        ON CONFLICT(club_id) DO UPDATE SET name=EXCLUDED.name,stadium=EXCLUDED.stadium,region_id=EXCLUDED.region_id,team_id=EXCLUDED.team_id,games_played=EXCLUDED.games_played,wins=EXCLUDED.wins,ties=EXCLUDED.ties,losses=EXCLUDED.losses,goals=EXCLUDED.goals,goals_against=EXCLUDED.goals_against,skill_rating=EXCLUDED.skill_rating,win_streak=EXCLUDED.win_streak,unbeaten_streak=EXCLUDED.unbeaten_streak,promotions=EXCLUDED.promotions,relegations=EXCLUDED.relegations,best_division=EXCLUDED.best_division,updated_at=EXCLUDED.updated_at""",(club_id,official,d.get("stadium"),str(d.get("regionId") or ""),str(d.get("teamId") or ""),o.get("gamesPlayed"),o.get("wins"),o.get("ties"),o.get("losses"),o.get("goals"),o.get("goalsAgainst"),o.get("skillRating"),o.get("winStreak"),o.get("unbeatenStreak"),o.get("promotions"),o.get("relegations"),o.get("bestDivision"),now))
        execq(c,"INSERT INTO user_clubs(user_id,club_id) VALUES(%s,%s) ON CONFLICT DO NOTHING",(user_id,club_id))
        execq(c,"DELETE FROM members WHERE club_id=%s",(club_id,))
        for m in members:
            execq(c,"""INSERT INTO members VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(club_id,name) DO UPDATE SET player_name=EXCLUDED.player_name,position=EXCLUDED.position,overall=EXCLUDED.overall,games_played=EXCLUDED.games_played,win_rate=EXCLUDED.win_rate,goals=EXCLUDED.goals,assists=EXCLUDED.assists,average_rating=EXCLUDED.average_rating,motm=EXCLUDED.motm,shot_success_rate=EXCLUDED.shot_success_rate,passes_made=EXCLUDED.passes_made,pass_success_rate=EXCLUDED.pass_success_rate,tackles_made=EXCLUDED.tackles_made,tackle_success_rate=EXCLUDED.tackle_success_rate,clean_sheets_def=EXCLUDED.clean_sheets_def,clean_sheets_gk=EXCLUDED.clean_sheets_gk,red_cards=EXCLUDED.red_cards,updated_at=EXCLUDED.updated_at""",(club_id,str(m.get("name") or m.get("playerName") or "Jogador"),str(m.get("playerName") or ""),str(m.get("position") or ""),m.get("overall"),m.get("gamesPlayed"),m.get("winRate"),m.get("goals"),m.get("assists"),m.get("averageRating"),m.get("manOfTheMatch"),m.get("shotSuccessRate"),m.get("passesMade"),m.get("passSuccessRate"),m.get("tacklesMade"),m.get("tackleSuccessRate"),m.get("cleanSheetsDef"),m.get("cleanSheetsGK"),m.get("redCards"),now))
    new_matches=0
    for mt in MATCH_TYPES:
        with conn() as c:
            for m in safe_matches(club_id,mt):
                mid=m["matchId"]
                if not mid:continue
                existed=qone(c,"SELECT 1 FROM matches WHERE club_id=%s AND match_id=%s",(club_id,mid))
                execq(c,"""INSERT INTO matches VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(club_id,match_id) DO UPDATE SET match_type=EXCLUDED.match_type,timestamp=EXCLUDED.timestamp,club_name=EXCLUDED.club_name,opponent_id=EXCLUDED.opponent_id,opponent_name=EXCLUDED.opponent_name,goals=EXCLUDED.goals,goals_against=EXCLUDED.goals_against,result=EXCLUDED.result,dnf=EXCLUDED.dnf""",(club_id,mid,mt,m["timestamp"],m.get("clubName") or official,m.get("opponentId"),m.get("opponentName"),m["goals"],m["goalsAgainst"],m["result"],int(bool(m["dnf"]))))
                if not existed:new_matches+=1
                players=((m.get("_raw") or {}).get("players") or {}).get(str(club_id)) or {}
                for pid,s in players.items():
                    if not isinstance(s,dict):continue
                    name=str(s.get("playername") or s.get("name") or pid or "Jogador")
                    def num(k):
                        try:return float(s.get(k,0) or 0)
                        except:return 0
                    vals=(club_id,mid,mt,m["timestamp"],str(pid),name,str(s.get("pos") or ""),num("rating"),int(num("goals")),int(num("assists")),int(num("shots")),int(num("passesmade")),int(num("passattempts")),int(num("tacklesmade")),int(num("tackleattempts")),int(num("saves")),int(num("cleansheetsany")),int(num("mom")),int(num("redcards")),int(num("secondsPlayed")))
                    execq(c,"""INSERT INTO player_matches VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(club_id,match_id,player_id,name) DO UPDATE SET rating=EXCLUDED.rating,goals=EXCLUDED.goals,assists=EXCLUDED.assists,shots=EXCLUDED.shots,passes_made=EXCLUDED.passes_made,pass_attempts=EXCLUDED.pass_attempts,tackles_made=EXCLUDED.tackles_made,tackle_attempts=EXCLUDED.tackle_attempts,saves=EXCLUDED.saves,clean_sheet=EXCLUDED.clean_sheet,motm=EXCLUDED.motm,red_cards=EXCLUDED.red_cards,seconds_played=EXCLUDED.seconds_played""",vals)
    with conn() as c:total=qone(c,"SELECT COUNT(*) n FROM matches WHERE club_id=%s",(club_id,))["n"]
    return {"club_id":club_id,"name":official,"new_matches":new_matches,"total_matches":total}

try:init_db()
except Exception as e:print("DB init:",e)

@app.get("/favicon.ico",include_in_schema=False)
def favicon():return Response(status_code=204)
@app.get("/",response_class=HTMLResponse)
def home(request:Request):return templates.TemplateResponse(request=request,name="index.html",context={})
@app.get("/health")
def health():
    try:
        with conn() as c:qone(c,"SELECT 1")
        return {"status":"ok","database":"postgresql"}
    except Exception as e:raise HTTPException(503,str(e))
@app.get("/api/me")
def me(request:Request):
    u=current_user(request);return {"authenticated":bool(u),"user":dict(u) if u else None}
@app.post("/api/login")
async def login(request:Request,response:Response):
    body=await request.json();username=str(body.get("username","")).strip().lower();password=str(body.get("password",""))
    with conn() as c:u=qone(c,"SELECT * FROM users WHERE lower(username)=%s AND active=TRUE",(username,))
    if not u or not verify_password(password,u["password_hash"]):raise HTTPException(401,"Usuário ou senha incorretos.")
    token=secrets.token_urlsafe(32)
    with conn() as c:execq(c,"INSERT INTO sessions(token,user_id) VALUES(%s,%s)",(token,u["id"]))
    response.set_cookie("fc27_session",token,httponly=True,secure=True,samesite="lax",max_age=2592000)
    return {"ok":True,"username":u["username"],"is_admin":u["is_admin"]}
@app.post("/api/logout")
def logout(request:Request,response:Response):
    tok=request.cookies.get("fc27_session")
    if tok:
        with conn() as c:execq(c,"DELETE FROM sessions WHERE token=%s",(tok,))
    response.delete_cookie("fc27_session");return {"ok":True}
@app.get("/api/admin/users")
def users(request:Request):
    require_admin(request)
    with conn() as c:return qall(c,"SELECT id,username,is_admin,active,created_at FROM users ORDER BY username")
@app.post("/api/admin/users")
async def create_user(request:Request):
    require_admin(request);b=await request.json();username=str(b.get("username","")).strip().lower();password=str(b.get("password",""))
    if len(username)<3 or len(password)<6:raise HTTPException(400,"Usuário mínimo 3 caracteres e senha mínimo 6.")
    try:
        with conn() as c:execq(c,"INSERT INTO users(username,password_hash) VALUES(%s,%s)",(username,hash_password(password)))
    except Exception:raise HTTPException(409,"Usuário já existe.")
    return {"ok":True}
@app.get("/api/clubs")
def clubs(request:Request):
    u=require_user(request)
    with conn() as c:
        if u["is_admin"]:return qall(c,"SELECT c.*,(SELECT COUNT(*) FROM matches m WHERE m.club_id=c.club_id) total_matches FROM clubs c ORDER BY updated_at DESC")
        return qall(c,"SELECT c.*,(SELECT COUNT(*) FROM matches m WHERE m.club_id=c.club_id) total_matches FROM clubs c JOIN user_clubs uc ON uc.club_id=c.club_id WHERE uc.user_id=%s ORDER BY c.updated_at DESC",(u["id"],))
@app.delete("/api/clubs/{club_id}")
def remove_club(request:Request,club_id:str):
    u=require_user(request)
    if u["is_admin"]: raise HTTPException(400,"O administrador não pode remover clubes pela lista global.")
    with conn() as c:
        linked=qone(c,"SELECT 1 FROM user_clubs WHERE user_id=%s AND club_id=%s",(u["id"],club_id))
        if not linked: raise HTTPException(404,"Este clube não está nos seus clubes.")
        execq(c,"DELETE FROM user_clubs WHERE user_id=%s AND club_id=%s",(u["id"],club_id))
    return {"ok":True,"club_id":club_id}
@app.post("/api/update")
def update(request:Request,q:str=Query(min_length=2)):
    u=require_user(request)
    try:return update_club(q,u["id"])
    except (FC27APIError,ValueError) as e:raise HTTPException(502,str(e))
    except Exception as e:raise HTTPException(500,f"Falha ao atualizar: {e}")
@app.get("/api/club/{club_id}")
def club(request:Request,club_id:str,limit:int=20):
    assert_club(request,club_id)
    with conn() as c:
        cr=qone(c,"SELECT * FROM clubs WHERE club_id=%s",(club_id,))
        if not cr:raise HTTPException(404,"Clube não salvo.")
        return {"club":cr,"members":qall(c,"SELECT * FROM members WHERE club_id=%s ORDER BY goals DESC,assists DESC,average_rating DESC",(club_id,)),"matches":qall(c,"SELECT * FROM matches WHERE club_id=%s ORDER BY timestamp DESC LIMIT %s",(club_id,max(1,min(limit,1000))))}
@app.get("/api/club/{club_id}/matches")
def matches(request:Request,club_id:str,limit:int=50,match_type:str="all"):
    assert_club(request,club_id);sql="SELECT * FROM matches WHERE club_id=%s";args=[club_id]
    if match_type!="all":sql+=" AND match_type=%s";args.append(match_type)
    sql+=" ORDER BY timestamp DESC LIMIT %s";args.append(max(1,min(limit,2000)))
    with conn() as c:return qall(c,sql,args)
@app.get("/api/club/{club_id}/match/{match_id}")
def match_detail(request:Request,club_id:str,match_id:str):
    assert_club(request,club_id)
    with conn() as c:
        m=qone(c,"SELECT * FROM matches WHERE club_id=%s AND match_id=%s",(club_id,match_id))
        if not m:raise HTTPException(404,"Partida não encontrada.")
        return {"match":m,"players":qall(c,"SELECT * FROM player_matches WHERE club_id=%s AND match_id=%s ORDER BY rating DESC,goals DESC,assists DESC",(club_id,match_id))}
@app.get("/api/club/{club_id}/rankings")
def rankings(request:Request,club_id:str,last:int=0):
    assert_club(request,club_id)
    with conn() as c:
        extra="";args=[club_id]
        if last>0:extra=" AND match_id IN (SELECT match_id FROM matches WHERE club_id=%s ORDER BY timestamp DESC LIMIT %s)";args+=[club_id,min(last,1000)]
        return qall(c,"SELECT name,COUNT(*) games,SUM(goals) goals,SUM(assists) assists,ROUND(AVG(rating)::numeric,2) average_rating,SUM(motm) motm FROM player_matches WHERE club_id=%s"+extra+" GROUP BY name ORDER BY goals DESC,assists DESC,average_rating DESC",args)
@app.get("/api/club/{club_id}/player/{name}")
def player(request:Request,club_id:str,name:str):
    assert_club(request,club_id)
    with conn() as c:r=qall(c,"SELECT p.*,m.opponent_name,m.goals club_goals,m.goals_against,m.result FROM player_matches p JOIN matches m ON m.club_id=p.club_id AND m.match_id=p.match_id WHERE p.club_id=%s AND p.name=%s ORDER BY p.timestamp DESC",(club_id,name))
    if not r:raise HTTPException(404,"Jogador não encontrado.")
    return r
@app.get("/api/club/{club_id}/export.xlsx")
def export_excel(request:Request,club_id:str):
    assert_club(request,club_id)
    with conn() as c:club=qone(c,"SELECT * FROM clubs WHERE club_id=%s",(club_id,));members=qall(c,"SELECT * FROM members WHERE club_id=%s",(club_id,));matches=qall(c,"SELECT * FROM matches WHERE club_id=%s ORDER BY timestamp DESC",(club_id,));stats=qall(c,"SELECT * FROM player_matches WHERE club_id=%s ORDER BY timestamp DESC",(club_id,))
    wb=Workbook();ws=wb.active;ws.title="Clube"
    def fill(ws,rows):
        if not rows:return
        hs=list(rows[0].keys());ws.append(hs)
        for cell in ws[1]:cell.fill=PatternFill("solid",fgColor="0B1730");cell.font=Font(color="FFFFFF",bold=True)
        for r in rows:ws.append([r.get(h) for h in hs])
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    fill(ws,[club]);fill(wb.create_sheet("Jogadores"),members);fill(wb.create_sheet("Partidas"),matches);fill(wb.create_sheet("Stats Jogadores"),stats)
    bio=io.BytesIO();wb.save(bio);bio.seek(0);fn="FC27_"+str(club["name"]).replace(" ","_")+".xlsx"
    return StreamingResponse(bio,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{fn}"'})
